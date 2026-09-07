from __future__ import annotations

import difflib
import io
import json
import os
import re
import shutil
import subprocess
import tempfile
import zipfile
from dataclasses import dataclass
from pathlib import Path
from typing import Callable

import fitz  # PyMuPDF
from PIL import Image, ImageChops, ImageEnhance
from docx import Document as DocxDocument
from openpyxl import Workbook
from pptx import Presentation
from pptx.util import Inches
from pypdf import PdfReader, PdfWriter


@dataclass
class ToolResult:
    path: Path
    filename: str
    media_type: str = "application/octet-stream"
    message: str = "Operazione completata"


class ToolError(RuntimeError):
    pass


def _find_executable(*names: str) -> str | None:
    for name in names:
        found = shutil.which(name)
        if found:
            return found
    if os.name == "nt":
        candidates = []
        pf = os.environ.get("PROGRAMFILES", r"C:\Program Files")
        pfx86 = os.environ.get("PROGRAMFILES(X86)", r"C:\Program Files (x86)")
        candidates += [
            Path(pf) / "LibreOffice" / "program" / "soffice.exe",
            Path(pfx86) / "LibreOffice" / "program" / "soffice.exe",
            Path(pf) / "Tesseract-OCR" / "tesseract.exe",
            Path(pfx86) / "Tesseract-OCR" / "tesseract.exe",
        ]
        # Ghostscript version folders vary.
        for base in [Path(pf) / "gs", Path(pfx86) / "gs"]:
            if base.exists():
                candidates += list(base.glob("*/bin/gswin64c.exe"))
                candidates += list(base.glob("*/bin/gswin32c.exe"))
        for candidate in candidates:
            if candidate.exists() and candidate.name.lower() in {n.lower() for n in names}:
                return str(candidate)
    return None


def dependency_status() -> dict:
    soffice = _find_executable("soffice", "libreoffice", "soffice.exe")
    tesseract = _find_executable("tesseract", "tesseract.exe")
    ghostscript = _find_executable("gs", "gswin64c.exe", "gswin32c.exe")
    ollama = _find_executable("ollama", "ollama.exe")
    return {
        "core": {"available": True, "label": "Motore PDF locale", "detail": f"PyMuPDF {fitz.VersionBind}"},
        "libreoffice": {"available": bool(soffice), "label": "LibreOffice", "detail": soffice or "Non rilevato"},
        "tesseract": {"available": bool(tesseract), "label": "Tesseract OCR", "detail": tesseract or "Non rilevato"},
        "ghostscript": {"available": bool(ghostscript), "label": "Ghostscript", "detail": ghostscript or "Non rilevato"},
        "ollama": {"available": bool(ollama), "label": "Ollama (AI locale)", "detail": ollama or "Non rilevato"},
    }


def _open_pdf(path: Path, password: str | None = None) -> fitz.Document:
    try:
        doc = fitz.open(path)
    except Exception as exc:
        raise ToolError(f"Impossibile aprire il PDF: {exc}") from exc
    if doc.needs_pass:
        if not password or not doc.authenticate(password):
            doc.close()
            raise ToolError("Il PDF è protetto: inserisci la password corretta.")
    return doc


def _save_optimized(doc: fitz.Document, path: Path) -> None:
    doc.save(path, garbage=4, clean=True, deflate=True, deflate_images=True, deflate_fonts=True)


def _parse_pages(spec: str, page_count: int) -> list[int]:
    """Parse 1-based specs: 1,3,5-8,10-end."""
    spec = (spec or "").strip().lower().replace(" ", "")
    if not spec or spec in {"all", "tutte"}:
        return list(range(page_count))
    result: list[int] = []
    for token in spec.split(","):
        if not token:
            continue
        if "-" in token:
            a, b = token.split("-", 1)
            start = 1 if a in {"", "start", "inizio"} else int(a)
            end = page_count if b in {"", "end", "fine"} else int(b)
            if start > end:
                start, end = end, start
            result.extend(range(start - 1, end))
        else:
            n = page_count if token in {"end", "fine"} else int(token)
            result.append(n - 1)
    unique: list[int] = []
    for p in result:
        if p < 0 or p >= page_count:
            raise ToolError(f"Pagina fuori intervallo: {p + 1}. Il documento ha {page_count} pagine.")
        if p not in unique:
            unique.append(p)
    if not unique:
        raise ToolError("Nessuna pagina valida selezionata.")
    return unique


def merge(files: list[Path], options: dict, outdir: Path) -> ToolResult:
    if len(files) < 2:
        raise ToolError("Seleziona almeno due PDF da unire.")
    result = fitz.open()
    password = options.get("password") or None
    for f in files:
        src = _open_pdf(f, password)
        result.insert_pdf(src)
        src.close()
    out = outdir / "giumag_unito.pdf"
    _save_optimized(result, out)
    result.close()
    return ToolResult(out, out.name, "application/pdf", f"Uniti {len(files)} PDF")


def split(files: list[Path], options: dict, outdir: Path) -> ToolResult:
    src = _open_pdf(files[0], options.get("password") or None)
    mode = options.get("mode", "ranges")
    outputs: list[Path] = []
    if mode == "every":
        ranges = [[i] for i in range(src.page_count)]
    else:
        raw = options.get("ranges", "1-end")
        ranges = []
        for part in raw.split(";"):
            ranges.append(_parse_pages(part.strip(), src.page_count))
    for idx, pages in enumerate(ranges, 1):
        doc = fitz.open()
        for p in pages:
            doc.insert_pdf(src, from_page=p, to_page=p)
        path = outdir / f"parte_{idx:02d}.pdf"
        _save_optimized(doc, path)
        doc.close()
        outputs.append(path)
    src.close()
    if len(outputs) == 1:
        return ToolResult(outputs[0], outputs[0].name, "application/pdf")
    return _zip_results(outputs, outdir / "giumag_diviso.zip", f"Create {len(outputs)} parti")


def extract_pages(files: list[Path], options: dict, outdir: Path) -> ToolResult:
    src = _open_pdf(files[0], options.get("password") or None)
    pages = _parse_pages(options.get("pages", "1"), src.page_count)
    separate = bool(options.get("separate", False))
    if separate:
        outputs = []
        for p in pages:
            doc = fitz.open()
            doc.insert_pdf(src, from_page=p, to_page=p)
            path = outdir / f"pagina_{p+1}.pdf"
            _save_optimized(doc, path)
            doc.close()
            outputs.append(path)
        src.close()
        return _zip_results(outputs, outdir / "pagine_estratte.zip", f"Estratte {len(outputs)} pagine")
    doc = fitz.open()
    for p in pages:
        doc.insert_pdf(src, from_page=p, to_page=p)
    out = outdir / "pagine_estratte.pdf"
    _save_optimized(doc, out)
    doc.close(); src.close()
    return ToolResult(out, out.name, "application/pdf", f"Estratte {len(pages)} pagine")


def remove_pages(files: list[Path], options: dict, outdir: Path) -> ToolResult:
    doc = _open_pdf(files[0], options.get("password") or None)
    remove = set(_parse_pages(options.get("pages", "1"), doc.page_count))
    keep = [i for i in range(doc.page_count) if i not in remove]
    if not keep:
        doc.close(); raise ToolError("Non puoi rimuovere tutte le pagine.")
    outdoc = fitz.open()
    for p in keep:
        outdoc.insert_pdf(doc, from_page=p, to_page=p)
    out = outdir / "pagine_rimosse.pdf"
    _save_optimized(outdoc, out)
    outdoc.close(); doc.close()
    return ToolResult(out, out.name, "application/pdf", f"Rimosse {len(remove)} pagine")


def reorder_pages(files: list[Path], options: dict, outdir: Path) -> ToolResult:
    doc = _open_pdf(files[0], options.get("password") or None)
    pages = _parse_pages(options.get("order", "1-end"), doc.page_count)
    if len(pages) != doc.page_count:
        doc.close(); raise ToolError("L'ordine deve includere tutte le pagine, una sola volta.")
    outdoc = fitz.open()
    for p in pages:
        outdoc.insert_pdf(doc, from_page=p, to_page=p)
    out = outdir / "pdf_riordinato.pdf"
    _save_optimized(outdoc, out)
    outdoc.close(); doc.close()
    return ToolResult(out, out.name, "application/pdf")


def rotate_pdf(files: list[Path], options: dict, outdir: Path) -> ToolResult:
    doc = _open_pdf(files[0], options.get("password") or None)
    angle = int(options.get("angle", 90)) % 360
    pages = _parse_pages(options.get("pages", "all"), doc.page_count)
    for p in pages:
        page = doc[p]
        page.set_rotation((page.rotation + angle) % 360)
    out = outdir / "pdf_ruotato.pdf"
    _save_optimized(doc, out); doc.close()
    return ToolResult(out, out.name, "application/pdf", f"Ruotate {len(pages)} pagine")


def compress_pdf(files: list[Path], options: dict, outdir: Path) -> ToolResult:
    level = options.get("level", "balanced")
    out = outdir / "pdf_compresso.pdf"
    gs = _find_executable("gs", "gswin64c.exe", "gswin32c.exe")
    if level in {"strong", "extreme"} and gs:
        setting = "/ebook" if level == "strong" else "/screen"
        cmd = [gs, "-sDEVICE=pdfwrite", "-dCompatibilityLevel=1.7", f"-dPDFSETTINGS={setting}",
               "-dNOPAUSE", "-dQUIET", "-dBATCH", f"-sOutputFile={out}", str(files[0])]
        proc = subprocess.run(cmd, capture_output=True, text=True)
        if proc.returncode == 0 and out.exists():
            return ToolResult(out, out.name, "application/pdf", "Compressione avanzata completata con Ghostscript")
    doc = _open_pdf(files[0], options.get("password") or None)
    _save_optimized(doc, out); doc.close()
    msg = "PDF ottimizzato localmente"
    if level in {"strong", "extreme"} and not gs:
        msg += " (Ghostscript non rilevato: usata compressione standard)"
    return ToolResult(out, out.name, "application/pdf", msg)


def repair_pdf(files: list[Path], options: dict, outdir: Path) -> ToolResult:
    out = outdir / "pdf_riparato.pdf"
    try:
        doc = _open_pdf(files[0], options.get("password") or None)
        _save_optimized(doc, out); doc.close()
        return ToolResult(out, out.name, "application/pdf", "Struttura PDF riscritta e ripulita")
    except Exception as first_exc:
        gs = _find_executable("gs", "gswin64c.exe", "gswin32c.exe")
        if not gs:
            raise ToolError(f"Riparazione non riuscita: {first_exc}")
        cmd = [gs, "-sDEVICE=pdfwrite", "-dNOPAUSE", "-dQUIET", "-dBATCH", f"-sOutputFile={out}", str(files[0])]
        proc = subprocess.run(cmd, capture_output=True, text=True)
        if proc.returncode != 0 or not out.exists():
            raise ToolError("Ghostscript non è riuscito a recuperare il documento.")
        return ToolResult(out, out.name, "application/pdf", "PDF recuperato tramite Ghostscript")


def watermark(files: list[Path], options: dict, outdir: Path) -> ToolResult:
    doc = _open_pdf(files[0], options.get("password") or None)
    text = str(options.get("text", "Giumag"))
    opacity = float(options.get("opacity", 0.18))
    size = float(options.get("size", 42))
    position = options.get("position", "center")
    for page in doc:
        rect = page.rect
        tw = fitz.get_text_length(text, fontname="helv", fontsize=size)
        if position == "top": x, y = max(24, (rect.width-tw)/2), 48
        elif position == "bottom": x, y = max(24, (rect.width-tw)/2), rect.height-30
        else: x, y = max(24, (rect.width-tw)/2), rect.height/2
        page.insert_text((x, y), text, fontsize=size, fontname="helv", fill_opacity=opacity, overlay=True)
    out = outdir / "pdf_filigrana.pdf"
    _save_optimized(doc, out); doc.close()
    return ToolResult(out, out.name, "application/pdf")


def page_numbers(files: list[Path], options: dict, outdir: Path) -> ToolResult:
    doc = _open_pdf(files[0], options.get("password") or None)
    start = int(options.get("start", 1))
    size = float(options.get("size", 11))
    position = options.get("position", "bottom-center")
    template = str(options.get("template", "{n}"))
    for i, page in enumerate(doc):
        text = template.replace("{n}", str(start+i)).replace("{total}", str(doc.page_count))
        width = fitz.get_text_length(text, fontname="helv", fontsize=size)
        if position == "bottom-left": x, y = 36, page.rect.height - 24
        elif position == "bottom-right": x, y = page.rect.width - width - 36, page.rect.height - 24
        elif position == "top-left": x, y = 36, 30
        elif position == "top-right": x, y = page.rect.width - width - 36, 30
        elif position == "top-center": x, y = (page.rect.width-width)/2, 30
        else: x, y = (page.rect.width-width)/2, page.rect.height - 24
        page.insert_text((x, y), text, fontsize=size, fontname="helv", overlay=True)
    out = outdir / "pdf_numerato.pdf"
    _save_optimized(doc, out); doc.close()
    return ToolResult(out, out.name, "application/pdf")


def crop_pdf(files: list[Path], options: dict, outdir: Path) -> ToolResult:
    doc = _open_pdf(files[0], options.get("password") or None)
    left = float(options.get("left", 0)); right = float(options.get("right", 0))
    top = float(options.get("top", 0)); bottom = float(options.get("bottom", 0))
    pages = _parse_pages(options.get("pages", "all"), doc.page_count)
    for p in pages:
        page = doc[p]
        r = page.cropbox
        nr = fitz.Rect(r.x0 + left, r.y0 + top, r.x1 - right, r.y1 - bottom)
        if nr.width <= 20 or nr.height <= 20:
            doc.close(); raise ToolError("I margini di ritaglio sono troppo grandi.")
        page.set_cropbox(nr)
    out = outdir / "pdf_ritagliato.pdf"
    _save_optimized(doc, out); doc.close()
    return ToolResult(out, out.name, "application/pdf")


def protect_pdf(files: list[Path], options: dict, outdir: Path) -> ToolResult:
    password = str(options.get("new_password", ""))
    if not password:
        raise ToolError("Inserisci una password.")
    reader = PdfReader(str(files[0]), password=options.get("password") or None)
    writer = PdfWriter()
    for p in reader.pages: writer.add_page(p)
    writer.encrypt(user_password=password, owner_password=password, algorithm="AES-256")
    out = outdir / "pdf_protetto.pdf"
    with out.open("wb") as fp: writer.write(fp)
    return ToolResult(out, out.name, "application/pdf", "PDF cifrato con AES-256")


def unlock_pdf(files: list[Path], options: dict, outdir: Path) -> ToolResult:
    password = str(options.get("password", ""))
    try:
        reader = PdfReader(str(files[0]), password=password or None)
        writer = PdfWriter()
        for p in reader.pages: writer.add_page(p)
    except Exception as exc:
        raise ToolError("Password non valida o PDF non decifrabile.") from exc
    out = outdir / "pdf_sbloccato.pdf"
    with out.open("wb") as fp: writer.write(fp)
    return ToolResult(out, out.name, "application/pdf")


def redact_text(files: list[Path], options: dict, outdir: Path) -> ToolResult:
    query = str(options.get("query", "")).strip()
    if not query: raise ToolError("Inserisci il testo da censurare.")
    doc = _open_pdf(files[0], options.get("password") or None)
    count = 0
    case_sensitive = bool(options.get("case_sensitive", False))
    terms = [q.strip() for q in query.split(";") if q.strip()]
    for page in doc:
        for term in terms:
            # PyMuPDF search is case-insensitive by default for ASCII; exact-case filtering requires word inspection.
            rects = page.search_for(term)
            if case_sensitive:
                words = page.get_text("words")
                rects = [fitz.Rect(w[:4]) for w in words if w[4] == term]
            for rect in rects:
                page.add_redact_annot(rect, fill=(0, 0, 0))
                count += 1
        if count:
            page.apply_redactions()
    out = outdir / "pdf_censurato.pdf"
    _save_optimized(doc, out); doc.close()
    return ToolResult(out, out.name, "application/pdf", f"Rimosse permanentemente {count} occorrenze")


def images_to_pdf(files: list[Path], options: dict, outdir: Path) -> ToolResult:
    if not files: raise ToolError("Seleziona almeno un'immagine.")
    margin = int(options.get("margin", 0))
    images: list[Image.Image] = []
    for path in files:
        try:
            im = Image.open(path)
            if getattr(im, "n_frames", 1) > 1:
                im.seek(0)
            im = im.convert("RGB")
            if margin > 0:
                canvas = Image.new("RGB", (im.width + margin*2, im.height + margin*2), "white")
                canvas.paste(im, (margin, margin)); im = canvas
            images.append(im)
        except Exception as exc:
            raise ToolError(f"Immagine non valida: {path.name}") from exc
    out = outdir / "immagini.pdf"
    images[0].save(out, "PDF", save_all=True, append_images=images[1:], resolution=options.get("dpi", 150))
    for im in images: im.close()
    return ToolResult(out, out.name, "application/pdf", f"Convertite {len(files)} immagini")


def pdf_to_images(files: list[Path], options: dict, outdir: Path) -> ToolResult:
    doc = _open_pdf(files[0], options.get("password") or None)
    fmt = options.get("format", "png").lower()
    dpi = int(options.get("dpi", 150))
    pages = _parse_pages(options.get("pages", "all"), doc.page_count)
    outputs = []
    for p in pages:
        pix = doc[p].get_pixmap(dpi=dpi, alpha=False)
        ext = "jpg" if fmt in {"jpg", "jpeg"} else "png"
        path = outdir / f"pagina_{p+1}.{ext}"
        if ext == "jpg": pix.save(str(path), jpg_quality=int(options.get("quality", 88)))
        else: pix.save(str(path))
        outputs.append(path)
    doc.close()
    return _zip_results(outputs, outdir / "pagine_immagini.zip", f"Esportate {len(outputs)} immagini")


def office_to_pdf(files: list[Path], options: dict, outdir: Path) -> ToolResult:
    soffice = _find_executable("soffice", "libreoffice", "soffice.exe")
    if not soffice:
        raise ToolError("LibreOffice non è installato/rilevato. Serve per convertire Word, Excel e PowerPoint in PDF.")
    outputs = []
    for f in files:
        cmd = [soffice, "--headless", "--convert-to", "pdf", "--outdir", str(outdir), str(f)]
        env = os.environ.copy(); env.setdefault("HOME", str(outdir))
        proc = subprocess.run(cmd, capture_output=True, text=True, env=env, timeout=120)
        out = outdir / f"{f.stem}.pdf"
        if proc.returncode != 0 or not out.exists():
            raise ToolError(f"LibreOffice non è riuscito a convertire {f.name}: {proc.stderr.strip() or proc.stdout.strip()}")
        outputs.append(out)
    if len(outputs) == 1:
        return ToolResult(outputs[0], outputs[0].name, "application/pdf", "Conversione eseguita con LibreOffice locale")
    return _zip_results(outputs, outdir / "documenti_pdf.zip", f"Convertiti {len(outputs)} documenti")


def pdf_to_word(files: list[Path], options: dict, outdir: Path) -> ToolResult:
    doc = _open_pdf(files[0], options.get("password") or None)
    word = DocxDocument()
    word.core_properties.title = files[0].stem
    for i, page in enumerate(doc):
        blocks = page.get_text("blocks", sort=True)
        for block in blocks:
            text = str(block[4]).strip()
            if text:
                p = word.add_paragraph()
                p.add_run(text)
        if i < doc.page_count - 1:
            word.add_page_break()
    doc.close()
    out = outdir / "pdf_convertito.docx"
    word.save(out)
    return ToolResult(out, out.name, "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                      "DOCX modificabile creato dal testo del PDF; layout complessi possono richiedere ritocchi")


def pdf_to_pptx(files: list[Path], options: dict, outdir: Path) -> ToolResult:
    doc = _open_pdf(files[0], options.get("password") or None)
    prs = Presentation()
    prs.slide_width = Inches(13.333); prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]
    for i, page in enumerate(doc):
        pix = page.get_pixmap(dpi=int(options.get("dpi", 144)), alpha=False)
        img = outdir / f"_slide_{i+1}.png"; pix.save(str(img))
        slide = prs.slides.add_slide(blank)
        # fit image preserving aspect ratio
        page_ratio = pix.width / pix.height
        slide_ratio = prs.slide_width / prs.slide_height
        if page_ratio > slide_ratio:
            w = prs.slide_width; h = int(w / page_ratio); x = 0; y = int((prs.slide_height-h)/2)
        else:
            h = prs.slide_height; w = int(h * page_ratio); y = 0; x = int((prs.slide_width-w)/2)
        slide.shapes.add_picture(str(img), x, y, width=w, height=h)
        img.unlink(missing_ok=True)
    # remove default first slide only if it was created (Presentation starts with 0 slides in python-pptx default)
    doc.close()
    out = outdir / "pdf_convertito.pptx"; prs.save(out)
    return ToolResult(out, out.name, "application/vnd.openxmlformats-officedocument.presentationml.presentation",
                      "PPTX creato con una slide fedele per ogni pagina PDF")


def pdf_to_excel(files: list[Path], options: dict, outdir: Path) -> ToolResult:
    import pdfplumber
    wb = Workbook(); wb.remove(wb.active)
    found_tables = 0
    with pdfplumber.open(str(files[0]), password=options.get("password") or None) as pdf:
        for i, page in enumerate(pdf.pages, 1):
            tables = page.extract_tables()
            if tables:
                for t_idx, table in enumerate(tables, 1):
                    ws = wb.create_sheet(f"P{i}_T{t_idx}"[:31])
                    for row in table:
                        ws.append(["" if cell is None else cell for cell in row])
                    found_tables += 1
            else:
                text = page.extract_text() or ""
                ws = wb.create_sheet(f"Pagina_{i}"[:31])
                for line in text.splitlines():
                    ws.append([line])
    if not wb.sheetnames:
        wb.create_sheet("Vuoto")
    out = outdir / "pdf_convertito.xlsx"; wb.save(out)
    msg = f"Estratte {found_tables} tabelle" if found_tables else "Nessuna tabella rilevata: testo inserito per pagina"
    return ToolResult(out, out.name, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", msg)


def pdf_to_markdown(files: list[Path], options: dict, outdir: Path) -> ToolResult:
    doc = _open_pdf(files[0], options.get("password") or None)
    chunks = [f"# {files[0].stem}\n"]
    for i, page in enumerate(doc, 1):
        chunks.append(f"\n## Pagina {i}\n")
        text = page.get_text("text", sort=True).strip()
        chunks.append(text if text else "_Nessun testo estraibile._")
    doc.close()
    out = outdir / "documento.md"; out.write_text("\n\n".join(chunks), encoding="utf-8")
    return ToolResult(out, out.name, "text/markdown", "Testo estratto in Markdown")


def ocr_pdf(files: list[Path], options: dict, outdir: Path) -> ToolResult:
    import pytesseract
    tess = _find_executable("tesseract", "tesseract.exe")
    if not tess:
        raise ToolError("Tesseract OCR non è installato/rilevato.")
    pytesseract.pytesseract.tesseract_cmd = tess
    doc = _open_pdf(files[0], options.get("password") or None)
    lang = options.get("lang", "ita+eng")
    dpi = int(options.get("dpi", 220))
    writer = PdfWriter()
    for page in doc:
        pix = page.get_pixmap(dpi=dpi, alpha=False)
        im = Image.open(io.BytesIO(pix.tobytes("png")))
        try:
            pdf_bytes = pytesseract.image_to_pdf_or_hocr(im, extension="pdf", lang=lang)
        except Exception as exc:
            doc.close()
            raise ToolError(f"OCR non riuscito. Verifica che le lingue '{lang}' siano installate in Tesseract: {exc}") from exc
        r = PdfReader(io.BytesIO(pdf_bytes))
        writer.add_page(r.pages[0])
        im.close()
    doc.close()
    out = outdir / "pdf_ocr.pdf"
    with out.open("wb") as fp: writer.write(fp)
    return ToolResult(out, out.name, "application/pdf", f"OCR locale completato ({lang})")


def compare_pdf(files: list[Path], options: dict, outdir: Path) -> ToolResult:
    if len(files) < 2: raise ToolError("Seleziona due PDF da confrontare.")
    a = _open_pdf(files[0], options.get("password") or None)
    b = _open_pdf(files[1], options.get("password2") or options.get("password") or None)
    text_a = "\n".join(p.get_text("text") for p in a)
    text_b = "\n".join(p.get_text("text") for p in b)
    diff = "\n".join(difflib.unified_diff(text_a.splitlines(), text_b.splitlines(), fromfile=files[0].name, tofile=files[1].name, lineterm=""))
    txt = outdir / "differenze_testo.txt"; txt.write_text(diff or "Nessuna differenza testuale rilevata.", encoding="utf-8")
    diff_pdf = fitz.open()
    max_pages = max(a.page_count, b.page_count)
    visual_changed = 0
    for i in range(max_pages):
        if i < a.page_count:
            pa = a[i].get_pixmap(dpi=100, alpha=False)
            ia = Image.open(io.BytesIO(pa.tobytes("png"))).convert("RGB")
        else:
            ia = Image.new("RGB", (1000, 1400), "white")
        if i < b.page_count:
            pb = b[i].get_pixmap(dpi=100, alpha=False)
            ib = Image.open(io.BytesIO(pb.tobytes("png"))).convert("RGB")
        else:
            ib = Image.new("RGB", ia.size, "white")
        w = min(ia.width, ib.width); h = min(ia.height, ib.height)
        ia = ia.resize((w,h)); ib = ib.resize((w,h))
        d = ImageChops.difference(ia, ib).convert("L")
        bbox = d.getbbox()
        if bbox: visual_changed += 1
        d = ImageEnhance.Contrast(d).enhance(3.0).convert("RGB")
        bio = io.BytesIO(); d.save(bio, format="PNG")
        page = diff_pdf.new_page(width=w*0.72, height=h*0.72)
        page.insert_image(page.rect, stream=bio.getvalue())
        ia.close(); ib.close(); d.close()
    visual = outdir / "differenze_visive.pdf"; _save_optimized(diff_pdf, visual); diff_pdf.close(); a.close(); b.close()
    return _zip_results([txt, visual], outdir / "confronto_pdf.zip", f"Confronto completato: {visual_changed} pagine con differenze visive")


def sign_visual(files: list[Path], options: dict, outdir: Path) -> ToolResult:
    pdfs = [f for f in files if f.suffix.lower() == ".pdf"]
    images = [f for f in files if f.suffix.lower() in {".png", ".jpg", ".jpeg", ".webp"}]
    if not pdfs or not images: raise ToolError("Seleziona un PDF e un'immagine PNG/JPG della firma.")
    doc = _open_pdf(pdfs[0], options.get("password") or None)
    page_num = int(options.get("page", doc.page_count)) - 1
    if page_num < 0 or page_num >= doc.page_count: doc.close(); raise ToolError("Pagina firma non valida.")
    page = doc[page_num]
    width = float(options.get("width", 150)); height = float(options.get("height", 60))
    pos = options.get("position", "bottom-right")
    margin = 36
    if pos == "bottom-left": x0, y0 = margin, page.rect.height-height-margin
    elif pos == "top-right": x0, y0 = page.rect.width-width-margin, margin
    elif pos == "top-left": x0, y0 = margin, margin
    elif pos == "center": x0, y0 = (page.rect.width-width)/2, (page.rect.height-height)/2
    else: x0, y0 = page.rect.width-width-margin, page.rect.height-height-margin
    rect = fitz.Rect(x0, y0, x0+width, y0+height)
    page.insert_image(rect, filename=str(images[0]), keep_proportion=True, overlay=True)
    out = outdir / "pdf_firmato_visivamente.pdf"; _save_optimized(doc, out); doc.close()
    return ToolResult(out, out.name, "application/pdf", "Firma visiva applicata (non è una firma digitale crittografica)")


def remove_metadata(files: list[Path], options: dict, outdir: Path) -> ToolResult:
    doc = _open_pdf(files[0], options.get("password") or None)
    doc.set_metadata({})
    try:
        doc.del_xml_metadata()
    except Exception:
        pass
    out = outdir / "pdf_senza_metadati.pdf"; _save_optimized(doc, out); doc.close()
    return ToolResult(out, out.name, "application/pdf", "Metadati principali rimossi")


def flatten_pdf(files: list[Path], options: dict, outdir: Path) -> ToolResult:
    doc = _open_pdf(files[0], options.get("password") or None)
    if hasattr(doc, "bake"):
        doc.bake(annots=True, widgets=True)
    else:
        # Fallback: rasterize, preserving appearance but not selectable text.
        rebuilt = fitz.open()
        for p in doc:
            pix = p.get_pixmap(dpi=150, alpha=False)
            np = rebuilt.new_page(width=p.rect.width, height=p.rect.height)
            np.insert_image(np.rect, stream=pix.tobytes("png"))
        doc.close(); doc = rebuilt
    out = outdir / "pdf_appiattito.pdf"; _save_optimized(doc, out); doc.close()
    return ToolResult(out, out.name, "application/pdf", "Annotazioni e moduli appiattiti")


def pdfa(files: list[Path], options: dict, outdir: Path) -> ToolResult:
    gs = _find_executable("gs", "gswin64c.exe", "gswin32c.exe")
    if not gs:
        raise ToolError("Ghostscript non è installato/rilevato. È necessario per la conversione PDF/A locale.")
    level = str(options.get("level", "2"))
    out = outdir / "documento_pdfa.pdf"
    cmd = [gs, "-dPDFA=" + level, "-dBATCH", "-dNOPAUSE", "-dNOOUTERSAVE", "-sDEVICE=pdfwrite",
           "-sColorConversionStrategy=UseDeviceIndependentColor", "-dPDFACompatibilityPolicy=1",
           f"-sOutputFile={out}", str(files[0])]
    proc = subprocess.run(cmd, capture_output=True, text=True, timeout=180)
    if proc.returncode != 0 or not out.exists():
        raise ToolError("Conversione PDF/A non riuscita con Ghostscript: " + (proc.stderr.strip() or proc.stdout.strip())[-500:])
    return ToolResult(out, out.name, "application/pdf", f"Conversione PDF/A-{level} completata con Ghostscript")


def local_ai(files: list[Path], options: dict, outdir: Path) -> ToolResult:
    ollama = _find_executable("ollama", "ollama.exe")
    if not ollama: raise ToolError("Ollama non è installato/rilevato. Installa Ollama e un modello locale per usare questa funzione.")
    doc = _open_pdf(files[0], options.get("password") or None)
    text = "\n".join(p.get_text("text") for p in doc); doc.close()
    if not text.strip(): raise ToolError("Il PDF non contiene testo estraibile. Esegui prima l'OCR.")
    action = options.get("action", "summarize")
    model = options.get("model", "llama3.2")
    if action == "translate":
        language = options.get("language", "italiano")
        instruction = f"Traduci fedelmente il seguente documento in {language}, mantenendo titoli e struttura quando possibile. Restituisci solo la traduzione."
    else:
        instruction = "Riassumi il documento seguente in italiano in modo chiaro, con una sintesi iniziale e punti chiave. Restituisci solo il riassunto."
    # Limit very large documents to avoid accidental runaway; user can chunk manually later.
    payload = (instruction + "\n\nDOCUMENTO:\n" + text[:120000])
    proc = subprocess.run([ollama, "run", model], input=payload, capture_output=True, text=True, timeout=600)
    if proc.returncode != 0: raise ToolError("Ollama ha restituito un errore: " + proc.stderr.strip()[-500:])
    out = outdir / ("traduzione_locale.md" if action == "translate" else "riassunto_locale.md")
    out.write_text(proc.stdout, encoding="utf-8")
    return ToolResult(out, out.name, "text/markdown", f"Elaborazione completata in locale con {model}")



def add_text(files: list[Path], options: dict, outdir: Path) -> ToolResult:
    doc = _open_pdf(files[0], options.get("password") or None)
    text = str(options.get("text", "")).strip()
    if not text:
        doc.close(); raise ToolError("Inserisci il testo da aggiungere.")
    page_num = int(options.get("page", 1)) - 1
    if page_num < 0 or page_num >= doc.page_count:
        doc.close(); raise ToolError("Pagina non valida.")
    page = doc[page_num]
    size = float(options.get("size", 14))
    position = options.get("position", "top-left")
    margin = 42
    box_w = min(page.rect.width * 0.75, 420)
    box_h = min(page.rect.height * 0.35, 240)
    if position == "top-right":
        x0, y0 = page.rect.width - box_w - margin, margin
    elif position == "bottom-left":
        x0, y0 = margin, page.rect.height - box_h - margin
    elif position == "bottom-right":
        x0, y0 = page.rect.width - box_w - margin, page.rect.height - box_h - margin
    elif position == "center":
        x0, y0 = (page.rect.width-box_w)/2, (page.rect.height-box_h)/2
    else:
        x0, y0 = margin, margin
    rect = fitz.Rect(x0, y0, x0+box_w, y0+box_h)
    page.insert_textbox(rect, text, fontsize=size, fontname="helv", overlay=True)
    out = outdir / "pdf_modificato.pdf"
    _save_optimized(doc, out); doc.close()
    return ToolResult(out, out.name, "application/pdf", "Testo aggiunto al PDF")


def fill_forms(files: list[Path], options: dict, outdir: Path) -> ToolResult:
    doc = _open_pdf(files[0], options.get("password") or None)
    raw = str(options.get("values", ""))
    mapping = {}
    for line in raw.splitlines():
        if "=" in line:
            k, v = line.split("=", 1)
            mapping[k.strip()] = v.strip()
    if not mapping:
        doc.close(); raise ToolError("Inserisci almeno un valore nel formato nome_campo=valore.")
    changed = 0
    available = []
    for page in doc:
        widgets = list(page.widgets() or [])
        for widget in widgets:
            if widget.field_name:
                available.append(widget.field_name)
            if widget.field_name in mapping:
                widget.field_value = mapping[widget.field_name]
                widget.update()
                changed += 1
    if changed == 0:
        names = ", ".join(dict.fromkeys(available)) or "nessun campo rilevato"
        doc.close(); raise ToolError("Nessun campo corrispondente. Campi disponibili: " + names)
    if bool(options.get("flatten", False)) and hasattr(doc, "bake"):
        doc.bake(annots=True, widgets=True)
    out = outdir / "modulo_compilato.pdf"
    _save_optimized(doc, out); doc.close()
    return ToolResult(out, out.name, "application/pdf", f"Compilati {changed} campi modulo")


def extract_assets(files: list[Path], options: dict, outdir: Path) -> ToolResult:
    doc = _open_pdf(files[0], options.get("password") or None)
    outputs = []
    text_path = outdir / "testo_estratto.txt"
    text_path.write_text("\n\n".join(f"--- Pagina {i+1} ---\n{p.get_text('text', sort=True)}" for i, p in enumerate(doc)), encoding="utf-8")
    outputs.append(text_path)
    seen = set()
    count = 0
    for page_no, page in enumerate(doc, 1):
        for img in page.get_images(full=True):
            xref = img[0]
            if xref in seen:
                continue
            seen.add(xref)
            try:
                info = doc.extract_image(xref)
                ext = info.get("ext", "bin")
                path = outdir / f"immagine_p{page_no}_{xref}.{ext}"
                path.write_bytes(info["image"])
                outputs.append(path); count += 1
            except Exception:
                pass
    doc.close()
    return _zip_results(outputs, outdir / "contenuti_pdf.zip", f"Estratti testo e {count} immagini")

def _zip_results(paths: list[Path], zip_path: Path, message: str) -> ToolResult:
    with zipfile.ZipFile(zip_path, "w", zipfile.ZIP_DEFLATED) as zf:
        for p in paths:
            zf.write(p, p.name)
    return ToolResult(zip_path, zip_path.name, "application/zip", message)


TOOL_HANDLERS: dict[str, Callable[[list[Path], dict, Path], ToolResult]] = {
    "merge": merge,
    "split": split,
    "extract": extract_pages,
    "remove-pages": remove_pages,
    "reorder": reorder_pages,
    "rotate": rotate_pdf,
    "compress": compress_pdf,
    "repair": repair_pdf,
    "watermark": watermark,
    "page-numbers": page_numbers,
    "crop": crop_pdf,
    "protect": protect_pdf,
    "unlock": unlock_pdf,
    "redact": redact_text,
    "images-to-pdf": images_to_pdf,
    "pdf-to-images": pdf_to_images,
    "office-to-pdf": office_to_pdf,
    "pdf-to-word": pdf_to_word,
    "pdf-to-pptx": pdf_to_pptx,
    "pdf-to-excel": pdf_to_excel,
    "pdf-to-markdown": pdf_to_markdown,
    "ocr": ocr_pdf,
    "compare": compare_pdf,
    "sign-visual": sign_visual,
    "remove-metadata": remove_metadata,
    "flatten": flatten_pdf,
    "pdfa": pdfa,
    "local-ai": local_ai,
    "add-text": add_text,
    "fill-forms": fill_forms,
    "extract-assets": extract_assets,
}
